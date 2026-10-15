"""Document generation service — Word (.docx) and Excel (.xlsx) export."""

from io import BytesIO

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter


def _style_docx(doc: Document) -> None:
    """Apply consistent styling to a Word document."""
    style = doc.styles["Normal"]
    font = style.font
    font.name = "SimSun"
    font.size = Pt(12)
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_after = Pt(6)

    for level in range(1, 4):
        heading = doc.styles[f"Heading {level}"]
        heading.font.name = "SimHei"
        heading.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)


def _markdown_to_docx_blocks(md_text: str) -> list[dict]:
    """Crude Markdown-to-blocks parser. Returns list of {type, text, level}."""
    blocks = []
    for line in md_text.split("\n"):
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("### "):
            blocks.append({"type": "h3", "text": stripped[4:], "level": 3})
        elif stripped.startswith("## "):
            blocks.append({"type": "h2", "text": stripped[3:], "level": 2})
        elif stripped.startswith("# "):
            blocks.append({"type": "h1", "text": stripped[2:], "level": 1})
        elif stripped.startswith("- ") or stripped.startswith("* "):
            blocks.append({"type": "bullet", "text": stripped[2:]})
        elif stripped.startswith("**") and "**" in stripped[2:]:
            blocks.append({"type": "bold", "text": stripped})
        else:
            blocks.append({"type": "p", "text": stripped})
    return blocks


def generate_docx(content: str, title: str = "调查报告") -> bytes:
    """Convert Markdown content into a styled .docx and return bytes."""
    doc = Document()
    _style_docx(doc)

    # Title
    t = doc.add_heading(title, level=1)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for block in _markdown_to_docx_blocks(content):
        if block["type"] in ("h1", "h2", "h3"):
            doc.add_heading(block["text"], level=block["level"])
        elif block["type"] == "bullet":
            doc.add_paragraph(block["text"], style="List Bullet")
        elif block["type"] == "bold":
            p = doc.add_paragraph()
            run = p.add_run(block["text"].strip("*"))
            run.bold = True
        else:
            doc.add_paragraph(block["text"])

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def generate_xlsx(content: str, title: str = "数据分析") -> bytes:
    """Convert structured text into a styled .xlsx workbook.

    If content contains a markdown table, it is parsed as rows.
    Otherwise the full text is placed in a single sheet as a report.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title[:31]

    header_font = Font(name="SimHei", size=12, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    cell_font = Font(name="SimSun", size=11)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    # Try to parse as markdown table
    lines = content.split("\n")
    table_lines = [l for l in lines if l.strip().startswith("|")]

    if len(table_lines) >= 2:
        header_cells = [c.strip() for c in table_lines[0].split("|") if c.strip()]
        for col_idx, h in enumerate(header_cells, 1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        row_idx = 2
        for line in table_lines[2:]:  # skip separator line
            cells = [c.strip() for c in line.split("|") if c.strip()]
            if not cells:
                continue
            for col_idx, val in enumerate(cells, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = cell_font
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.border = thin_border
            row_idx += 1

        for col_idx in range(1, len(header_cells) + 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = 20
    else:
        ws.merge_cells("A1:D1")
        title_cell = ws.cell(row=1, column=1, value=title)
        title_cell.font = Font(name="SimHei", size=16, bold=True)
        title_cell.alignment = Alignment(horizontal="center")

        for i, line in enumerate(lines, 3):
            if line.strip():
                ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=4)
                ws.cell(row=i, column=1, value=line.strip()).font = cell_font

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
